"""Сквозной тест печатных документов: предпросмотр A4 на телефоне, PDF через печать Chromium, выгрузка в Word.

PDF проверяется pypdf (число страниц, ориентация, текст), Word — python-docx (абзацы, таблицы).

    pip install pypdf python-docx
    python tests/e2e/print_flow.py [--shots каталог]
"""
import io
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import docx  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402
from pypdf import PdfReader  # noqa: E402

from common import PHONE, call, check_layout, fixture, server, shot  # noqa: E402


def pdf_of(page):
    page.emulate_media(media="print")
    data = page.pdf(prefer_css_page_size=True, print_background=True)
    page.emulate_media(media="screen")
    return PdfReader(io.BytesIO(data))


def main():
    with server() as (base, key):
        status, created = call(base, "POST", "/api/v1/events", {"program": fixture("rwp-2026.program.json")}, {"X-API-Key": key})
        assert status == 201, created
        ev_id, slug = created["event"]["id"], created["event"]["slug"]
        token = created["invite_url"].rsplit("/", 1)[1]
        status, claimed = call(base, "POST", f"/api/v1/claim/{token}", {"email": "org@example.com", "password": "org-pass-1234"})
        auth = {"Authorization": claimed["token"]}
        program = call(base, "GET", f"/api/v1/events/{ev_id}/program", headers=auth)[1]["program"]
        day0_sessions = len(program["days"][0]["sessions"])

        with sync_playwright() as p:
            browser = p.chromium.launch()
            ctx = browser.new_context(**PHONE, accept_downloads=True)
            ctx.add_init_script("localStorage.setItem('conf_auth', JSON.stringify(%s))" % json.dumps({"token": claimed["token"], "record": claimed["record"]}))
            page = ctx.new_page()
            errors = []
            page.on("pageerror", lambda e: errors.append(str(e)))

            # из своего мероприятия — в документы; лист A4 помещается в 320 px
            page.goto(f"{base}/#/my/{ev_id}")
            page.get_by_role("link", name="Документы для печати").click()
            page.locator(".pd-program").wait_for()
            check_layout(page, "печать 320")
            box = page.locator(".print-scale").bounding_box()
            assert box["width"] <= 320, box
            shot(page, "30-print-program-320")

            # программа: PDF на нескольких страницах A4, текст как в официальной программе
            pdf = pdf_of(page)
            assert len(pdf.pages) >= 3, len(pdf.pages)
            mb = pdf.pages[0].mediabox
            assert float(mb.height) > float(mb.width), "книжная"
            text = "".join(pg.extract_text() for pg in pdf.pages)
            assert "07.10.2026 [Среда]" in text
            assert "Докладчик (онлайн): Е. С. Беленя" in re.sub(r"\s+", " ", text)

            # Word — программа
            with page.expect_download() as dl:
                page.get_by_role("button", name="Word").click()
            d = docx.Document(dl.value.path())
            paras = [x.text for x in d.paragraphs]
            assert any("Докладчик (онлайн): Е. С. Беленя" in t for t in paras)
            assert any(t.startswith("07.10.2026 [Среда]") for t in paras)

            # таблички на дверь за первый день: по заседанию на альбомный лист
            page.get_by_role("tab", name="На дверь").click()
            page.locator(".pd-door").first.wait_for()
            page.get_by_label("День").select_option("0")
            assert page.locator(".pd-door").count() == day0_sessions
            shot(page, "31-print-doors-320")
            pdf = pdf_of(page)
            assert len(pdf.pages) == day0_sessions, (len(pdf.pages), day0_sessions)
            mb = pdf.pages[0].mediabox
            assert float(mb.width) > float(mb.height), "альбомная"

            # протоколы: таблица с шапкой, Word читается python-docx
            page.get_by_role("tab", name="Протоколы").click()
            page.locator(".pd-protocol").first.wait_for()
            check_layout(page, "протоколы 320")
            shot(page, "32-print-protocols-320")
            with page.expect_download() as dl:
                page.get_by_role("button", name="Word").click()
            d = docx.Document(dl.value.path())
            assert len(d.tables) == page.locator(".pd-protocol").count()
            assert [c.text for c in d.tables[0].rows[0].cells] == ["№", "ФИО докладчика", "Тема доклада", "Место работы", "Форма участия", "Примечание"]

            # сертификаты
            page.get_by_role("tab", name="Сертификаты").click()
            page.locator(".pd-cert").first.wait_for()
            shot(page, "33-print-cert-320")

            # участники: версия для печати опубликованной программы без входа
            assert call(base, "POST", f"/api/v1/events/{ev_id}/publish", {}, auth)[0] == 200
            anon = browser.new_context(**PHONE).new_page()
            anon.goto(f"{base}/#/e/{slug}")
            anon.get_by_role("link", name="Версия для печати").click()
            anon.locator(".pd-program").wait_for()

            assert not errors, errors
            browser.close()
        print("e2e print: OK")


if __name__ == "__main__":
    main()
